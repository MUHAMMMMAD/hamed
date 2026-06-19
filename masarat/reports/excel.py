"""
تصدير ملف Excel احترافي (Cost Sheet) من نتيجة التحليل.
أوراق: ملخص تنفيذي · تسعير البنود · تفصيل التكلفة · المخاطر · التدفقات · النماذج.
"""
from __future__ import annotations

from pathlib import Path

from .. import config
from ..models import TenderAnalysis

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except Exception:  # pragma: no cover
    Workbook = None


_HEADER_FILL = "1F4E78"
_ACCENT_FILL = "D9E1F2"


def _style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 4, 50)


def export_excel(a: TenderAnalysis, path: str | Path | None = None) -> Path:
    if Workbook is None:
        raise RuntimeError("openpyxl غير مثبّت — ثبّته عبر: pip install openpyxl")

    path = Path(path) if path else config.OUTPUT_DIR / f"{a.tender.id}_cost_sheet.xlsx"
    wb = Workbook()

    # ----- 1) ملخص تنفيذي -----
    ws = wb.active
    ws.title = "ملخص تنفيذي"
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "MASARAT AI — تقرير التسعير التنفيذي"
    ws["A1"].font = Font(bold=True, size=14, color=_HEADER_FILL)
    rows = [
        ("المشروع", a.tender.name),
        ("العميل", a.tender.client),
        ("الموقع", a.tender.location),
        ("المدة (شهر)", a.tender.duration_months),
        ("القرار", a.decision.decision.value),
        ("درجة القرار", a.decision.score),
        ("احتمال الفوز %", a.decision.win_probability),
        ("مستوى المخاطر", a.risk.overall_level.value),
        ("التكلفة المباشرة", a.cost.direct_cost),
        ("إجمالي التكلفة", a.cost.total_cost),
        ("سعر البيع (قبل الضريبة)", a.cost.selling_price),
        ("السعر النهائي (شامل الضريبة)", a.cost.final_price),
        ("هامش الربح %", a.cost.margin_pct),
        ("أقصى تمويل مطلوب", a.cashflow.peak_funding_required),
    ]
    r = 3
    for label, value in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=_ACCENT_FILL)
        ws.cell(row=r, column=2, value=value)
        r += 1
    _autosize(ws)

    # ----- 2) تسعير البنود -----
    ws = wb.create_sheet("تسعير البنود")
    ws.sheet_view.rightToLeft = True
    headers = ["الكود", "الوصف", "الوحدة", "الكمية", "سعر الوحدة",
               "مواد", "عمالة", "معدات", "مقاول باطن", "الإجمالي المباشر"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for it in a.tender.items:
        ws.append([
            it.id, it.description, it.unit, it.quantity, it.unit_rate,
            it.material_cost, it.labor_cost, it.equipment_cost,
            it.subcontractor_cost, it.direct_total,
        ])
    ws.append(["", "", "", "", "", "", "", "", "الإجمالي", a.cost.direct_cost])
    ws.cell(row=ws.max_row, column=10).font = Font(bold=True)
    _autosize(ws)

    # ----- 3) تفصيل التكلفة -----
    ws = wb.create_sheet("تفصيل التكلفة")
    ws.sheet_view.rightToLeft = True
    ws.append(["البند", "القيمة (ريال)"])
    _style_header(ws, 1, 2)
    cost_rows = [
        ("المواد", a.cost.materials),
        ("العمالة", a.cost.labor),
        ("المعدات", a.cost.equipment),
        ("مقاولو الباطن", a.cost.subcontractors),
        ("التكلفة المباشرة", a.cost.direct_cost),
        ("مصاريف إدارية", a.cost.overhead),
        ("احتياطي", a.cost.contingency),
        ("تعديل المخاطر", a.cost.risk_adjustment),
        ("إجمالي التكلفة", a.cost.total_cost),
        ("الربح", a.cost.profit),
        ("سعر البيع (قبل الضريبة)", a.cost.selling_price),
        ("ضريبة القيمة المضافة", a.cost.vat),
        ("السعر النهائي (شامل الضريبة)", a.cost.final_price),
    ]
    for label, value in cost_rows:
        ws.append([label, value])
    _autosize(ws)

    # ----- 4) المخاطر -----
    ws = wb.create_sheet("المخاطر")
    ws.sheet_view.rightToLeft = True
    ws.append(["الفئة", "الوصف", "الدرجة", "المستوى", "التخفيف"])
    _style_header(ws, 1, 5)
    for ri in a.risk.items:
        ws.append([ri.category, ri.description, ri.score, ri.level.value, ri.mitigation])
    _autosize(ws)

    # ----- 5) التدفقات النقدية -----
    ws = wb.create_sheet("التدفقات النقدية")
    ws.sheet_view.rightToLeft = True
    ws.append(["الشهر", "التكلفة", "الإيراد", "الصافي", "التراكمي"])
    _style_header(ws, 1, 5)
    for m in a.cashflow.months:
        ws.append([m.month, m.cost, m.revenue, m.net, m.cumulative])
    _autosize(ws)

    # ----- 6) مجلس النماذج -----
    ws = wb.create_sheet("مجلس النماذج")
    ws.sheet_view.rightToLeft = True
    ws.append(["النموذج", "الدور", "الوضع", "الثقة %"])
    _style_header(ws, 1, 4)
    for key, ins in a.council_insights.items():
        ws.append([
            ins.get("display_name"),
            ins.get("role"),
            "تجريبي" if ins.get("offline") else "حقيقي",
            ins.get("confidence"),
        ])
    _autosize(ws)

    wb.save(path)
    return path
