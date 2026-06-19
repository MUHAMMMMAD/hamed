"""
توليد القوالب الجاهزة:
- قالب جدول كميات (BOQ) لرفع مشاريع جديدة.
- قالب قواعد الأسعار (مُعبّأ بالبيانات الحالية) لإدخال بيانات الشركة الفعلية.
"""
from __future__ import annotations

from pathlib import Path

from .. import config
from ..database import Database, db as default_db

# الحقول لكل ورقة في قالب الأسعار (ترويسات = أسماء الحقول لضمان استيراد دقيق)
PRICE_SHEETS = {
    "materials": ["id", "name_ar", "name_en", "unit", "unit_price", "category", "trend_pct", "last_updated"],
    "labor": ["id", "name_ar", "name_en", "daily_rate"],
    "equipment": ["id", "name_ar", "name_en", "unit", "rate"],
    "suppliers": ["id", "name", "materials", "lead_time_days", "price_stability", "rating", "last_updated"],
    "subcontractors": ["id", "name", "activity", "unit_price", "unit", "quality", "timeliness", "risk"],
    "projects": ["id", "name", "client", "type", "location", "contract_value",
                 "duration_months", "actual_cost", "actual_profit", "deviation_pct", "deviation_reason"],
}


def _style(ws, ncols: int) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.sheet_view.rightToLeft = True
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")


def generate_boq_template(path: str | Path | None = None) -> Path:
    """قالب جدول كميات فارغ مع أمثلة توضيحية."""
    from openpyxl import Workbook

    path = Path(path) if path else config.OUTPUT_DIR / "BOQ_template.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    headers = ["الكود", "الوصف", "الوحدة", "الكمية", "سعر الوحدة",
               "مواد", "عمالة", "معدات", "مقاول باطن"]
    ws.append(headers)
    _style(ws, len(headers))
    # أمثلة (سعر الوحدة إجمالي، أو فصّل المكوّنات لكل وحدة)
    ws.append(["B-01", "خرسانة مسلحة C30", "م3", 800, 342, "", "", "", ""])
    ws.append(["B-02", "حديد التسليح", "طن", 95, "", 3160, 1740, "", ""])
    ws.append(["B-03", "أعمال الكهرباء", "م2", 3500, 120, "", "", "", 120])
    note = ws.cell(
        row=ws.max_row + 2, column=1,
        value="ملاحظة: املأ 'سعر الوحدة' (إجمالي لكل وحدة)، أو فصّل المكوّنات "
              "(مواد/عمالة/معدات/مقاول باطن) لكل وحدة. الأعمدة الإضافية اختيارية.",
    )
    for col in ws.columns:
        first = col[0]
        ws.column_dimensions[first.column_letter].width = 18
    wb.save(path)
    return path


def generate_prices_template(
    path: str | Path | None = None, database: Database | None = None
) -> Path:
    """قالب قواعد الأسعار مُعبّأ بالبيانات الحالية ليعدّله المستخدم ببياناته."""
    from openpyxl import Workbook

    db = database or default_db
    path = Path(path) if path else config.OUTPUT_DIR / "PRICES_template.xlsx"
    wb = Workbook()
    wb.remove(wb.active)

    data_sources = {
        "materials": db.materials.values(),
        "labor": db.labor.values(),
        "equipment": db.equipment.values(),
        "suppliers": db.suppliers.values(),
        "subcontractors": db.subcontractors.values(),
        "projects": db.projects.values(),
    }

    for sheet, fields in PRICE_SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append(fields)
        _style(ws, len(fields))
        for obj in data_sources[sheet]:
            d = obj.model_dump()
            row = []
            for f in fields:
                val = d.get(f)
                if isinstance(val, list):           # suppliers.materials
                    val = ";".join(val)
                elif hasattr(val, "value"):         # enums (type/risk)
                    val = val.value
                row.append(val)
            ws.append(row)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 16

    wb.save(path)
    return path
