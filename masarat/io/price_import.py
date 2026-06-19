"""
استيراد قواعد الأسعار الفعلية من ملف Excel (قالب الأسعار) إلى النظام.
يكتب البيانات إلى ملفات data/seed/*.json ويعيد تحميل قاعدة البيانات فوراً.
بهذا تُدخل شركة مسارات بياناتها الحقيقية (مواد/عمالة/موردون/مشاريع سابقة).
"""
from __future__ import annotations

import json
from pathlib import Path

from .. import config
from ..database import db as default_db
from ..models import (
    Equipment,
    Labor,
    Material,
    PastProject,
    Subcontractor,
    Supplier,
)

SHEET_MODELS = {
    "materials": (Material, "materials.json"),
    "labor": (Labor, "labor.json"),
    "equipment": (Equipment, "equipment.json"),
    "suppliers": (Supplier, "suppliers.json"),
    "subcontractors": (Subcontractor, "subcontractors.json"),
    "projects": (PastProject, "projects.json"),
}


def import_prices(path: str | Path) -> dict:
    """يستورد كل أوراق قالب الأسعار ويعيد عدد السجلات لكل ورقة."""
    from openpyxl import load_workbook

    wb = load_workbook(Path(path), data_only=True)
    counts: dict[str, int] = {}
    errors: list[str] = []

    for sheet, (Model, filename) in SHEET_MODELS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]

        objs = []
        for r in rows[1:]:
            if r is None or all(c is None for c in r):
                continue
            raw = {
                headers[i]: r[i]
                for i in range(min(len(headers), len(r)))
                if headers[i]
            }
            if not raw.get("id"):
                continue
            # حقل المواد لدى الموردين: نص مفصول بفواصل منقوطة -> قائمة
            if sheet == "suppliers":
                mats = raw.get("materials")
                if isinstance(mats, str):
                    raw["materials"] = [m.strip() for m in mats.split(";") if m.strip()]
                elif mats is None:
                    raw["materials"] = []
            try:
                objs.append(Model(**raw))
            except Exception as exc:  # تخطّي الصفوف غير الصالحة مع تسجيلها
                errors.append(f"{sheet}:{raw.get('id')} -> {exc}")
                continue

        out = config.DATA_DIR / filename
        out.write_text(
            json.dumps([o.model_dump() for o in objs], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        counts[sheet] = len(objs)

    default_db.reload()   # تفعيل البيانات الجديدة فوراً
    result = {"imported": counts}
    if errors:
        result["errors"] = errors
    return result
