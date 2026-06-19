"""
قارئ جداول الكميات التلقائي (BOQ Auto-Reader).
يفهم ملفات Excel (.xlsx) و CSV بترويسات عربية أو إنجليزية، يكتشف الأعمدة
آلياً (الكود/الوصف/الوحدة/الكمية/سعر الوحدة...)، ويحوّلها إلى كائن Tender جاهز
للتسعير. يدعم جداول مُسعّرة مسبقاً أو فارغة الأسعار.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import List, Optional

from ..models import BOQItem, ProjectType, Tender

# كلمات مفتاحية لاكتشاف الأعمدة (عربي + إنجليزي)
COLUMN_KEYWORDS = {
    "code": ["الكود", "كود", "رقم البند", "رقم", "بند رقم", "item no", "code", "s/n", "sr", "م", "#", "ت", "no"],
    "description": ["الوصف", "وصف", "البند", "بيان", "الأعمال", "الاعمال", "description", "desc"],
    "unit": ["الوحدة", "وحدة", "unit", "uom"],
    "quantity": ["الكمية", "كمية", "كميه", "العدد", "quantity", "qty"],
    "unit_rate": ["سعر الوحدة", "السعر", "سعر", "الفئة", "فئة", "unit rate", "unit price", "rate", "price"],
    "total": ["الإجمالي", "الاجمالي", "اجمالي", "القيمة", "قيمة", "total", "amount"],
    "material": ["مواد", "خامات", "material"],
    "labor": ["عمالة", "عمال", "labor", "labour"],
    "equipment": ["معدات", "آليات", "اليات", "equipment"],
    "subcontractor": ["مقاول باطن", "باطن", "subcontractor", "sub-contractor"],
}

# ترتيب إسناد الأعمدة (الأكثر تحديداً أولاً لتجنّب التعارض)
# ملاحظة: unit_rate قبل unit حتى لا يلتقط عمودُ "الوحدة" خانةَ "سعر الوحدة".
ASSIGN_ORDER = [
    "code", "description", "quantity",
    "unit_rate", "total", "material", "labor", "equipment", "subcontractor", "unit",
]


def _norm(text) -> str:
    return str(text).strip().lower() if text is not None else ""


def _to_float(val) -> float:
    """تحويل قيمة (قد تحوي فواصل/عملة/نص) إلى رقم."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace(",", "").replace("٬", "")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group()) if m else 0.0


def _match_field(cell: str, field: str) -> bool:
    cell = _norm(cell)
    if not cell:
        return False
    for kw in COLUMN_KEYWORDS[field]:
        kw = kw.lower()
        if len(kw) <= 2:
            if cell == kw:
                return True
        elif kw in cell:
            return True
    return False


def _detect_header(row: List) -> dict:
    """يُرجع خريطة {field: col_index} للأعمدة المكتشفة في صف الترويسة."""
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for field in ASSIGN_ORDER:
        for idx, cell in enumerate(row):
            if idx in used:
                continue
            if _match_field(cell, field):
                mapping[field] = idx
                used.add(idx)
                break
    return mapping


def _find_header_row(rows: List[List]) -> tuple[int, dict]:
    """يبحث عن صف الترويسة ضمن أول 15 صفاً (الأعلى تطابقاً)."""
    best_idx, best_map, best_score = -1, {}, 0
    for i, row in enumerate(rows[:15]):
        mapping = _detect_header(row)
        score = sum(1 for k in ("description", "quantity", "unit") if k in mapping)
        if score > best_score:
            best_idx, best_map, best_score = i, mapping, score
    if "description" not in best_map or "quantity" not in best_map:
        raise ValueError(
            "تعذّر اكتشاف أعمدة الجدول. تأكد من وجود عمودي 'الوصف' و'الكمية' "
            "(أو استخدم القالب الجاهز)."
        )
    return best_idx, best_map


def _rows_from_excel(path: Path) -> List[List]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _rows_from_csv(path: Path) -> List[List]:
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return [row for row in csv.reader(f)]


def import_boq(
    path: str | Path,
    name: Optional[str] = None,
    client: str = "غير محدد",
    project_type: str = "private",
    location: str = "الرياض",
    duration_months: int = 12,
    tender_id: Optional[str] = None,
) -> Tender:
    """يقرأ ملف BOQ (Excel/CSV) ويُرجع مناقصة جاهزة للتحليل."""
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        rows = _rows_from_excel(path)
    elif path.suffix.lower() == ".csv":
        rows = _rows_from_csv(path)
    else:
        raise ValueError(f"صيغة غير مدعومة: {path.suffix} (المدعوم: .xlsx, .csv)")

    if not rows:
        raise ValueError("الملف فارغ")

    header_idx, col = _find_header_row(rows)
    items: List[BOQItem] = []

    for n, row in enumerate(rows[header_idx + 1:], start=1):
        def cell(field):
            i = col.get(field)
            return row[i] if i is not None and i < len(row) else None

        desc = _norm(cell("description"))
        qty = _to_float(cell("quantity"))
        if not desc or qty <= 0:          # تخطّي العناوين والصفوف الفارغة
            continue

        item = BOQItem(
            id=str(cell("code") or f"I-{n:03d}").strip(),
            description=str(cell("description")).strip(),
            unit=str(cell("unit") or "وحدة").strip(),
            quantity=qty,
            source="imported",
        )

        # تحديد سعر الوحدة من الأعمدة المتاحة (حسب الأولوية)
        comps = {
            "material": _to_float(cell("material")),
            "labor": _to_float(cell("labor")),
            "equipment": _to_float(cell("equipment")),
            "subcontractor": _to_float(cell("subcontractor")),
        }
        comp_sum = sum(comps.values())
        rate = _to_float(cell("unit_rate"))
        total = _to_float(cell("total"))

        if comp_sum > 0:                  # مكوّنات لكل وحدة
            item.unit_rate = round(comp_sum, 2)
            item.material_cost = round(comps["material"] * qty, 2)
            item.labor_cost = round(comps["labor"] * qty, 2)
            item.equipment_cost = round(comps["equipment"] * qty, 2)
            item.subcontractor_cost = round(comps["subcontractor"] * qty, 2)
        elif rate > 0:                    # سعر وحدة إجمالي
            item.unit_rate = round(rate, 2)
        elif total > 0:                   # قيمة سطر إجمالية
            item.unit_rate = round(total / qty, 2)
        # else: بند بلا سعر (سيظهر بصفر ويُعلَّم)

        items.append(item)

    if not items:
        raise ValueError("لم يتم العثور على بنود صالحة (وصف + كمية موجبة).")

    return Tender(
        id=tender_id or f"TND-{path.stem}",
        name=name or path.stem,
        client=client,
        type=ProjectType(project_type),
        location=location,
        duration_months=duration_months,
        items=items,
    )
